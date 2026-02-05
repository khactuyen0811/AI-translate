import torch
from docx import Document
from transformers import AutoTokenizer, AutoModelForCausalLM, BitsAndBytesConfig
from huggingface_hub import login
import os
import re
from google.colab import files

# ==========================================
# 1. CẤU HÌNH
# ==========================================
HF_TOKEN = "nhập token trên hugging face"
MODEL_ID = "google/gemma-2b-it"
# nhập tên file input và output
INPUT_WORD = "WSC2024_TP39_MC_actual_en.docx" 
OUTPUT_WORD = "Ban_Dich_test.docx"

def clean_ai_output(text, original):
    """Lọc bỏ lỗi từ chối dịch và xóa câu nói leo"""
    # Nếu AI trả về câu từ chối hoặc lỗi, lấy lại bản gốc tiếng Anh
    refusal_keywords = [
        "cannot", "unable", "not provided", "context", 
        "provide the text", "language model", "I am a"
    ]
    if any(word in text.lower() for word in refusal_keywords):
        return original

    # Xóa các cụm từ 'nói leo' phổ biến
    patterns = [
        r"Sure,.*:", r"Translation.*:", r"Here is.*:", 
        r"\*\*Translation:\*\*", r"\*\*Tiếng Việt:\*\*", 
        r"The translated text is:", r"Vietnamese translation:"
    ]
    for p in patterns:
        text = re.sub(p, "", text, flags=re.IGNORECASE)
    
    final_result = text.strip()
    # Nếu kết quả rỗng, trả về bản gốc
    return final_result if final_result else original

def should_skip(text):
    """Né các thanh căn lề (dấu chấm dài), mục lục và dòng số"""
    t = text.strip()
    if "...." in t or "____" in t: return True
    if t.isdigit(): return True
    if len(t) <= 1: return True
    return False

def main():
    login(HF_TOKEN)
    
    # Nạp model 4-bit để tối ưu bộ nhớ GPU T4
    q_config = BitsAndBytesConfig(load_in_4bit=True, bnb_4bit_compute_dtype=torch.float16)
    print("⏳ Đang khởi động AI Gemma (GPU T4)...")
    tokenizer = AutoTokenizer.from_pretrained(MODEL_ID)
    model = AutoModelForCausalLM.from_pretrained(MODEL_ID, quantization_config=q_config, device_map="auto")

    if not os.path.exists(INPUT_WORD):
        print(f"LỖI: Không tìm thấy file {INPUT_WORD}. Hãy upload file vào thư mục bên trái Colab!")
        return

    doc = Document(INPUT_WORD)
    all_items = []

    # Gom Paragraphs thông thường
    for para in doc.paragraphs:
        if para.text.strip(): all_items.append(para)

    # Gom nội dung trong Bảng
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if para.text.strip(): all_items.append(para)

    total = len(all_items)
    print(f"Bắt đầu dịch {total} mục.'")

    for i, item in enumerate(all_items):
        original_text = item.text.strip()
        
        # Kiểm tra xem có nên bỏ qua (mục lục/số trang) không
        if should_skip(original_text):
            continue

        # Prompt ép AI: Dịch hoặc chép lại, KHÔNG giải thích
        prompt = f"<start_of_turn>user\nTranslate to Vietnamese. If technical, keep English terms. If unable, repeat the original. Output ONLY translation.\nText: {original_text}<end_of_turn>\n<start_of_turn>model\n"
        
        inputs = tokenizer(prompt, return_tensors="pt").to("cuda")
        with torch.inference_mode():
            outputs = model.generate(**inputs, max_new_tokens=512, do_sample=False)
        
        decoded = tokenizer.decode(outputs[0][len(inputs['input_ids'][0]):], skip_special_tokens=True)
        
        # Dùng hàm lọc mới có truyền original_text để chống lỗi từ chối
        final_text = clean_ai_output(decoded, original_text)

        # Ghi đè chữ nhưng giữ định dạng căn lề (Alignment)
        old_alignment = item.alignment
        for run in item.runs:
            run.text = ""
        
        if item.runs:
            item.runs[0].text = final_text
        else:
            item.add_run(final_text)
        
        item.alignment = old_alignment

        # Cập nhật tiến độ mỗi 10 mục
        if (i+1) % 10 == 0 or (i+1) == total:
            print(f"Đã xong {i+1}/{total} mục ({(i+1)/total*100:.1f}%)", end="\r")
            doc.save(OUTPUT_WORD)

    doc.save(OUTPUT_WORD)
    print(f"\n HOÀN TẤT! Đang chuẩn bị tải file về máy...")
    files.download(OUTPUT_WORD)

if __name__ == "__main__":
    main()
